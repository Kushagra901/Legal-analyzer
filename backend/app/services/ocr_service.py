# ocr_service.py
"""
OCR and Multi-Format Document Parsing Service.
Handles image-to-text extraction, native PDF extraction, and text extraction
for DOCX, HTML, RTF, TXT, DOC, and images with fallback mechanisms.
"""
import io
import os
import re

import docx
import fitz  # PyMuPDF
import pytesseract
from bs4 import BeautifulSoup
from PIL import Image
from pypdf import PdfReader
from striprtf.striprtf import rtf_to_text


class OCRService:
    """
    Service class responsible for OCR, native parsing, and multi-format document text extraction.
    """

    def __init__(self) -> None:
        """
        Initializes OCRService, configuring the tesseract executable path for Windows.
        """
        self.tesseract_cmd = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        if os.path.exists(self.tesseract_cmd):
            pytesseract.pytesseract.tesseract_cmd = self.tesseract_cmd

    def should_use_ocr(self, native_text: str) -> bool:
        """
        Determines if OCR should be used based on characters extracted natively.
        If native text extraction yields fewer than 200 characters, we fall back to OCR.
        """
        cleaned = native_text.strip() if native_text else ""
        return len(cleaned) < 200

    def detect_format(self, file_source: str | bytes) -> str:
        """
        Heuristically detects the document format from magic bytes signature or file extension.
        Returns one of: 'pdf', 'docx', 'doc', 'rtf', 'html', 'txt', 'image', 'unknown'.
        """
        ext = ""
        header = b""
        if isinstance(file_source, str):
            _, ext = os.path.splitext(file_source.lower())
            ext = ext.lstrip(".")
            try:
                with open(file_source, "rb") as f:
                    header = f.read(16)
            except Exception:
                header = b""
        else:
            header = file_source[:16] if file_source else b""

        # 1. Signature checks
        if header.startswith(b"%PDF"):
            return "pdf"
        elif header.startswith(b"PK\x03\x04"):
            if ext == "docx" or not ext:
                return "docx"
        elif header.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
            return "doc"
        elif header.startswith(b"{\\rtf"):
            return "rtf"
        elif (
            header.startswith(b"\x89PNG\r\n\x1a\n")
            or header.startswith(b"\xff\xd8\xff")
            or header.startswith(b"II*\x00")
            or header.startswith(b"MM\x00*")
        ):
            return "image"

        # 2. HTML text check
        try:
            decoded = header.decode("utf-8", errors="ignore").strip().lower()
            if decoded.startswith("<html") or decoded.startswith("<!doctype html") or decoded.startswith("<head"):
                return "html"
        except Exception:
            pass

        # 3. Extension fallback
        if ext in ("pdf", "docx", "doc", "rtf", "html", "htm", "txt", "png", "jpg", "jpeg", "tiff", "tif"):
            if ext in ("jpg", "jpeg", "png", "tiff", "tif"):
                return "image"
            if ext in ("html", "htm"):
                return "html"
            return ext

        # 4. Text decode check
        if isinstance(file_source, bytes):
            try:
                file_source.decode("utf-8")
                return "txt"
            except UnicodeDecodeError:
                try:
                    file_source.decode("latin-1")
                    return "txt"
                except Exception:
                    pass

        return "unknown"

    def extract_text_natively(self, file_source: str | bytes) -> str:
        """
        Extract text from a PDF file using pypdf reader.
        """
        try:
            if isinstance(file_source, bytes):
                reader = PdfReader(io.BytesIO(file_source))
            else:
                reader = PdfReader(file_source)

            text_parts = []
            for page in reader.pages:
                text_content = page.extract_text()
                if text_content:
                    text_parts.append(text_content)
            return "\n".join(text_parts)
        except Exception as e:
            print(f"Error during native text extraction: {e}")
            return ""

    def extract_text_via_ocr(self, file_source: str | bytes) -> str:
        """
        Perform OCR on PDF pages by rendering them to images using PyMuPDF (fitz)
        and running Tesseract OCR.
        """
        try:
            if isinstance(file_source, bytes):
                doc = fitz.open(stream=file_source, filetype="pdf")
            else:
                doc = fitz.open(file_source)

            ocr_parts = []
            for page in doc:
                pix = page.get_pixmap()
                img_bytes = pix.tobytes("png")
                img = Image.open(io.BytesIO(img_bytes))
                text = pytesseract.image_to_string(img)
                if text:
                    ocr_parts.append(text)
            return "\n".join(ocr_parts)
        except Exception as e:
            print(f"Error during OCR text extraction: {e}")
            return ""

    def extract_docx(self, file_source: str | bytes) -> str:
        """
        Extract text from a DOCX document (including tables).
        """
        try:
            if isinstance(file_source, bytes):
                doc = docx.Document(io.BytesIO(file_source))
            else:
                doc = docx.Document(file_source)
            text_parts = []
            for para in doc.paragraphs:
                text_parts.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    text_parts.append(" | ".join([cell.text for cell in row.cells]))
            return "\n".join(text_parts)
        except Exception as e:
            print(f"Error during DOCX extraction: {e}")
            return ""

    def extract_html(self, file_source: str | bytes) -> str:
        """
        Extract text from an HTML document.
        """
        try:
            if isinstance(file_source, bytes):
                content = file_source.decode("utf-8", errors="ignore")
            else:
                with open(file_source, encoding="utf-8", errors="ignore") as f:
                    content = f.read()
            soup = BeautifulSoup(content, "html.parser")
            return soup.get_text(separator="\n")
        except Exception as e:
            print(f"Error during HTML extraction: {e}")
            return ""

    def extract_rtf(self, file_source: str | bytes) -> str:
        """
        Extract text from an RTF document.
        """
        try:
            if isinstance(file_source, bytes):
                content = file_source.decode("latin-1", errors="ignore")
            else:
                with open(file_source, encoding="latin-1", errors="ignore") as f:
                    content = f.read()
            return rtf_to_text(content)
        except Exception as e:
            print(f"Error during RTF extraction: {e}")
            return ""

    def extract_txt(self, file_source: str | bytes) -> str:
        """
        Extract text from a TXT document.
        """
        try:
            if isinstance(file_source, bytes):
                try:
                    return file_source.decode("utf-8")
                except UnicodeDecodeError:
                    return file_source.decode("latin-1", errors="ignore")
            else:
                with open(file_source, encoding="utf-8", errors="ignore") as f:
                    return f.read()
        except Exception as e:
            print(f"Error during TXT extraction: {e}")
            return ""

    def extract_doc_fallback(self, file_source: str | bytes) -> str:
        """
        Extract text heuristically from old binary DOC file.
        """
        try:
            if isinstance(file_source, str):
                with open(file_source, "rb") as f:
                    data = f.read()
            else:
                data = file_source
            pattern = re.compile(rb'[a-zA-Z0-9\s\.,;:!\?\-\'\"]{4,}')
            matches = pattern.findall(data)
            return "\n".join([m.decode("latin-1", errors="ignore") for m in matches])
        except Exception as e:
            print(f"Error during DOC fallback extraction: {e}")
            return ""

    def extract_image(self, file_source: str | bytes) -> str:
        """
        Extract text from image via pytesseract OCR.
        """
        try:
            if isinstance(file_source, bytes):
                img = Image.open(io.BytesIO(file_source))
            else:
                img = Image.open(file_source)
            return pytesseract.image_to_string(img)
        except Exception as e:
            print(f"Error during image OCR extraction: {e}")
            return ""

    def process_document(self, file_source: str | bytes) -> tuple[str, str, float]:
        """
        Detects the file type and uses the best extraction method.
        Returns:
            tuple[str, str, float]: (extracted_text, method_used, parsing_confidence)
        """
        file_format = self.detect_format(file_source)

        if file_format == "pdf":
            native_text = self.extract_text_natively(file_source)
            if self.should_use_ocr(native_text):
                print("Native PDF extraction yields low text. Running OCR fallback...")
                ocr_text = self.extract_text_via_ocr(file_source)
                return ocr_text, "ocr", 0.85
            else:
                return native_text, "native", 1.0

        elif file_format == "docx":
            text = self.extract_docx(file_source)
            confidence = 1.0 if len(text.strip()) >= 50 else 0.5
            return text, "native", confidence

        elif file_format == "doc":
            text = self.extract_doc_fallback(file_source)
            confidence = 0.5 if len(text.strip()) >= 50 else 0.2
            return text, "fallback", confidence

        elif file_format == "rtf":
            text = self.extract_rtf(file_source)
            confidence = 0.95 if len(text.strip()) >= 50 else 0.5
            return text, "native", confidence

        elif file_format == "html":
            text = self.extract_html(file_source)
            confidence = 1.0 if len(text.strip()) >= 50 else 0.5
            return text, "native", confidence

        elif file_format == "txt":
            text = self.extract_txt(file_source)
            confidence = 1.0 if len(text.strip()) >= 50 else 0.5
            return text, "native", confidence

        elif file_format == "image":
            text = self.extract_image(file_source)
            confidence = 0.85 if len(text.strip()) >= 50 else 0.4
            return text, "ocr", confidence

        else:
            # Try plain text decode as final fallback
            text = self.extract_txt(file_source)
            if len(text.strip()) >= 50:
                return text, "native", 0.5
            else:
                # String extract
                text = self.extract_doc_fallback(file_source)
                confidence = 0.3 if len(text.strip()) >= 50 else 0.1
                return text, "fallback", confidence

    def extract_text(self, file_path: str) -> str:
        """
        Deprecated. Legacy compatibility helper.
        """
        text_content, _, _ = self.process_document(file_path)
        return text_content
