"""
test_report_export.py
Unit and Integration tests for PDF and DOCX report generation, storage persistence,
and /report export endpoints using sample NDA contract data.
"""
import io
import os
import sys
import uuid
from unittest.mock import MagicMock

import docx
import fitz  # PyMuPDF
import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from app.core.auth import get_current_user
from app.core.database import get_db
from app.main import app
from app.models.database_models import (
    Clause,
    ComplianceCheck,
    Document,
    LegalReference,
    Report,
    RiskFlag,
    User,
)
from app.services.report_generator_service import ReportGeneratorService


@pytest.fixture
def mock_user():
    return User(
        id=uuid.uuid4(),
        email="attorney@example.com",
        role="user"
    )


@pytest.fixture
def sample_nda_report_data():
    """
    Sample structured findings data matching the sample NDA test document.
    """
    return {
        "document_id": str(uuid.uuid4()),
        "filename": "sample-nda-test-document.pdf",
        "generated_at": "August 27, 2026 at 14:00 UTC",
        "safety_score": 25,
        "risk_level": "HIGH",
        "document_overview": (
            "This document is a Mutual Non-Disclosure Agreement (NDA) between Acme Innovations Pvt. Ltd. "
            "and Jane Doe governing confidential disclosures during exploratory partnership discussions."
        ),
        "parties": [
            "Acme Innovations Pvt. Ltd.",
            "Jane Doe"
        ],
        "key_dates": {
            "effective_date": "Not specified",
            "expiration_date": "Until terminated upon written notice",
            "notice_period": "90 days prior written notice"
        },
        "clauses": [
            {
                "id": str(uuid.uuid4()),
                "clause_type": "Indemnification",
                "category": "Liability & Risk",
                "severity": "HIGH",
                "confidence_score": 0.98,
                "explanation": "Receiving Party is subject to uncapped unilateral indemnity obligations.",
                "clause_text": "The Receiving Party shall indemnify, defend, and hold harmless the Disclosing Party from and against any and all claims, losses, and damages."
            },
            {
                "id": str(uuid.uuid4()),
                "clause_type": "Termination",
                "category": "Term & Termination",
                "severity": "HIGH",
                "confidence_score": 0.95,
                "explanation": "Termination requires express written consent of Disclosing Party with 90 days notice.",
                "clause_text": "The Receiving Party may terminate this Agreement only upon ninety (90) days' prior written notice and with the express written consent of the Disclosing Party."
            },
            {
                "id": str(uuid.uuid4()),
                "clause_type": "Governing Law & Jurisdiction",
                "category": "Dispute Resolution & Jurisdiction",
                "severity": "LOW",
                "confidence_score": 0.99,
                "explanation": "Governed by the laws of India with exclusive jurisdiction in New Delhi.",
                "clause_text": "This Agreement shall be governed by and construed in accordance with the laws of India. The parties agree to submit to the exclusive jurisdiction of the courts located in New Delhi."
            }
        ],
        "risk_flags": [
            {
                "clause_type": "Indemnification",
                "severity": "HIGH",
                "explanation": "Uncapped indemnity creates disproportionate exposure for the receiving party."
            },
            {
                "clause_type": "Termination",
                "severity": "HIGH",
                "explanation": "Requirement for counterparty consent to terminate restricts standard exit rights."
            }
        ],
        "missing_sections": [
            "Force Majeure Provision",
            "Data Protection / GDPR Compliance Clause",
            "Return or Destruction of Confidential Materials",
            "Standard Exclusions from Confidential Information Definition"
        ],
        "plain_english_summary": (
            "This NDA heavily favors Acme Innovations. You are agreeing to keep their disclosures confidential "
            "with unlimited financial liability for breaches. Furthermore, you cannot terminate the agreement without "
            "their written permission, creating a high-risk asymmetrical obligation."
        ),
        "recommendations": [
            "Require bilateral mutual indemnification with a defined liability cap.",
            "Remove the requirement for counterparty written consent to terminate.",
            "Add standard exclusions for public domain information and legally compelled disclosures."
        ],
        "citations": [
            {
                "source": "Indian Contract Act, 1872",
                "citation": "Section 73 (Compensation for loss or damage caused by breach of contract)"
            }
        ]
    }


def test_generate_pdf_from_sample_nda(sample_nda_report_data):
    """
    Test generating publication-grade PDF report from sample NDA structured findings.
    """
    service = ReportGeneratorService()
    pdf_bytes = service.generate_pdf(sample_nda_report_data)

    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 2000
    assert pdf_bytes.startswith(b"%PDF")

    # Verify PyMuPDF can parse the generated PDF and extract text from all sections
    pdf_doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    assert pdf_doc.page_count >= 1

    full_text = ""
    for page in pdf_doc:
        full_text += page.get_text()

    assert "LEGAL REVIEW MEMORANDUM" in full_text
    assert "1.0 Document Overview" in full_text
    assert "2.0 Parties Involved" in full_text
    assert "3.0 Key Dates" in full_text
    assert "4.0 Important Clauses" in full_text
    assert "5.0 Risk Flags" in full_text
    assert "6.0 Missing or Unclear Provisions" in full_text
    assert "7.0 Plain-English Executive Summary" in full_text
    assert "8.0 Suggested Next Review Actions" in full_text

    # Verify key NDA entities
    assert "Acme Innovations Pvt. Ltd." in full_text
    assert "Jane Doe" in full_text
    assert "HIGH RISK" in full_text or "HIGH" in full_text


def test_generate_docx_from_sample_nda(sample_nda_report_data):
    """
    Test generating professional DOCX report from sample NDA structured findings.
    """
    service = ReportGeneratorService()
    docx_bytes = service.generate_docx(sample_nda_report_data)

    assert isinstance(docx_bytes, bytes)
    assert len(docx_bytes) > 2000
    assert docx_bytes.startswith(b"PK\x03\x04")

    # Verify python-docx can parse the generated document
    doc = docx.Document(io.BytesIO(docx_bytes))
    paragraphs_text = " ".join([p.text for p in doc.paragraphs])

    assert "LEGAL REVIEW MEMORANDUM" in paragraphs_text
    assert "1.0 Document Overview" in paragraphs_text
    assert "2.0 Parties Involved" in paragraphs_text
    assert "3.0 Key Dates & Timeline" in paragraphs_text
    assert "4.0 Important Clauses (Grouped by Category)" in paragraphs_text
    assert "5.0 Risk Flags & Warning Items" in paragraphs_text
    assert "6.0 Missing or Unclear Provisions" in paragraphs_text
    assert "7.0 Plain-English Executive Summary" in paragraphs_text
    assert "8.0 Suggested Next Review Actions" in paragraphs_text


def test_report_service_handles_sparse_data():
    """
    Test that ReportGeneratorService handles minimal or empty input gracefully.
    """
    service = ReportGeneratorService()
    minimal_data = {
        "document_id": str(uuid.uuid4()),
        "filename": "sparse.pdf",
        "safety_score": 100,
        "risk_level": "LOW"
    }

    pdf_bytes = service.generate_pdf(minimal_data)
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")

    docx_bytes = service.generate_docx(minimal_data)
    assert isinstance(docx_bytes, bytes)
    assert docx_bytes.startswith(b"PK\x03\x04")


def test_export_report_pdf_endpoint(mock_user, sample_nda_report_data):
    """
    Test POST /api/v1/reports/{document_id}/export endpoint for PDF format.
    """
    app.dependency_overrides[get_current_user] = lambda: mock_user

    mock_db = MagicMock(spec=Session)

    doc_uuid = uuid.UUID(sample_nda_report_data["document_id"])
    mock_doc = Document(
        id=doc_uuid,
        user_id=mock_user.id,
        filename="sample-nda-test-document.pdf",
        summary=sample_nda_report_data["document_overview"],
        safety_score=sample_nda_report_data["safety_score"],
        risk_level=sample_nda_report_data["risk_level"],
        document_overview=sample_nda_report_data["document_overview"],
        parties=sample_nda_report_data["parties"],
        key_dates=sample_nda_report_data["key_dates"],
        missing_sections=sample_nda_report_data["missing_sections"],
        plain_english_summary=sample_nda_report_data["plain_english_summary"],
        status="completed"
    )

    # Setup query mock
    def mock_query(model):
        query_mock = MagicMock()
        if model == Document:
            query_mock.filter.return_value.first.return_value = mock_doc
        elif model == Clause:
            query_mock.filter.return_value.all.return_value = []
        elif model == LegalReference:
            query_mock.filter.return_value.all.return_value = []
        elif model == ComplianceCheck:
            query_mock.filter.return_value.first.return_value = None
        elif model == Report:
            query_mock.filter.return_value.first.return_value = None
        elif model == RiskFlag:
            query_mock.filter.return_value.all.return_value = []
        return query_mock

    mock_db.query.side_effect = mock_query
    app.dependency_overrides[get_db] = lambda: mock_db

    client = TestClient(app)
    response = client.post(f"/api/v1/reports/{str(doc_uuid)}/export?format=pdf")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    res_json = response.json()
    assert res_json["document_id"] == str(doc_uuid)
    assert res_json["format"] == "pdf"
    assert "report_url" in res_json
    assert res_json["document_overview"] == sample_nda_report_data["document_overview"]
    assert res_json["parties"] == sample_nda_report_data["parties"]


def test_export_report_docx_endpoint(mock_user, sample_nda_report_data):
    """
    Test POST /api/v1/reports/{document_id}/export endpoint for DOCX format.
    """
    app.dependency_overrides[get_current_user] = lambda: mock_user

    mock_db = MagicMock(spec=Session)
    doc_uuid = uuid.UUID(sample_nda_report_data["document_id"])
    mock_doc = Document(
        id=doc_uuid,
        user_id=mock_user.id,
        filename="sample-nda-test-document.pdf",
        summary=sample_nda_report_data["document_overview"],
        safety_score=sample_nda_report_data["safety_score"],
        risk_level=sample_nda_report_data["risk_level"],
        document_overview=sample_nda_report_data["document_overview"],
        parties=sample_nda_report_data["parties"],
        key_dates=sample_nda_report_data["key_dates"],
        missing_sections=sample_nda_report_data["missing_sections"],
        plain_english_summary=sample_nda_report_data["plain_english_summary"],
        status="completed"
    )

    def mock_query(model):
        query_mock = MagicMock()
        if model == Document:
            query_mock.filter.return_value.first.return_value = mock_doc
        elif model == Clause:
            query_mock.filter.return_value.all.return_value = []
        elif model == LegalReference:
            query_mock.filter.return_value.all.return_value = []
        elif model == ComplianceCheck:
            query_mock.filter.return_value.first.return_value = None
        elif model == Report:
            query_mock.filter.return_value.first.return_value = None
        elif model == RiskFlag:
            query_mock.filter.return_value.all.return_value = []
        return query_mock

    mock_db.query.side_effect = mock_query
    app.dependency_overrides[get_db] = lambda: mock_db

    client = TestClient(app)
    response = client.post(f"/api/v1/reports/{str(doc_uuid)}/export?format=docx")

    app.dependency_overrides.clear()

    assert response.status_code == 200
    res_json = response.json()
    assert res_json["document_id"] == str(doc_uuid)
    assert res_json["format"] == "docx"
    assert "report_url" in res_json


def test_download_report_endpoints(mock_user, sample_nda_report_data):
    """
    Test GET /api/v1/reports/{document_id}/download for both PDF and DOCX.
    """
    app.dependency_overrides[get_current_user] = lambda: mock_user

    mock_db = MagicMock(spec=Session)
    doc_uuid = uuid.UUID(sample_nda_report_data["document_id"])
    mock_doc = Document(
        id=doc_uuid,
        user_id=mock_user.id,
        filename="sample-nda-test-document.pdf",
        summary=sample_nda_report_data["document_overview"],
        safety_score=sample_nda_report_data["safety_score"],
        risk_level=sample_nda_report_data["risk_level"],
        document_overview=sample_nda_report_data["document_overview"],
        parties=sample_nda_report_data["parties"],
        key_dates=sample_nda_report_data["key_dates"],
        missing_sections=sample_nda_report_data["missing_sections"],
        plain_english_summary=sample_nda_report_data["plain_english_summary"],
        status="completed"
    )

    def mock_query(model):
        query_mock = MagicMock()
        if model == Document:
            query_mock.filter.return_value.first.return_value = mock_doc
        elif model == Clause:
            query_mock.filter.return_value.all.return_value = []
        elif model == LegalReference:
            query_mock.filter.return_value.all.return_value = []
        elif model == ComplianceCheck:
            query_mock.filter.return_value.first.return_value = None
        elif model == Report:
            query_mock.filter.return_value.first.return_value = None
        elif model == RiskFlag:
            query_mock.filter.return_value.all.return_value = []
        return query_mock

    mock_db.query.side_effect = mock_query
    app.dependency_overrides[get_db] = lambda: mock_db

    client = TestClient(app)

    # Test PDF download
    pdf_res = client.get(f"/api/v1/reports/{str(doc_uuid)}/download?format=pdf")
    assert pdf_res.status_code == 200
    assert pdf_res.headers["content-type"] == "application/pdf"
    assert pdf_res.content.startswith(b"%PDF")
    assert pdf_res.headers["content-disposition"] == f'attachment; filename="Legal_Review_{str(doc_uuid)}.pdf"'

    # Test DOCX download
    docx_res = client.get(f"/api/v1/reports/{str(doc_uuid)}/download?format=docx")
    assert docx_res.status_code == 200
    assert "wordprocessingml" in docx_res.headers["content-type"]
    assert docx_res.content.startswith(b"PK\x03\x04")
    assert docx_res.headers["content-disposition"] == f'attachment; filename="Legal_Review_{str(doc_uuid)}.docx"'

    app.dependency_overrides.clear()


def test_document_run_report_endpoint(mock_user, sample_nda_report_data):
    """
    Test POST /api/v1/documents/{document_id}/report for both formats.
    """
    app.dependency_overrides[get_current_user] = lambda: mock_user

    mock_db = MagicMock(spec=Session)
    doc_uuid = uuid.UUID(sample_nda_report_data["document_id"])
    mock_doc = Document(
        id=doc_uuid,
        user_id=mock_user.id,
        filename="sample-nda-test-document.pdf",
        summary=sample_nda_report_data["document_overview"],
        safety_score=sample_nda_report_data["safety_score"],
        risk_level=sample_nda_report_data["risk_level"],
        document_overview=sample_nda_report_data["document_overview"],
        parties=sample_nda_report_data["parties"],
        key_dates=sample_nda_report_data["key_dates"],
        missing_sections=sample_nda_report_data["missing_sections"],
        plain_english_summary=sample_nda_report_data["plain_english_summary"],
        status="completed"
    )

    def mock_query(model):
        query_mock = MagicMock()
        if model == Document:
            query_mock.filter.return_value.first.return_value = mock_doc
        elif model == Clause:
            query_mock.filter.return_value.all.return_value = []
        elif model == LegalReference:
            query_mock.filter.return_value.all.return_value = []
        elif model == ComplianceCheck:
            query_mock.filter.return_value.first.return_value = None
        elif model == Report:
            query_mock.filter.return_value.first.return_value = None
        elif model == RiskFlag:
            query_mock.filter.return_value.all.return_value = []
        return query_mock

    mock_db.query.side_effect = mock_query
    app.dependency_overrides[get_db] = lambda: mock_db

    client = TestClient(app)

    # Test run_report with format=pdf
    res_pdf = client.post(f"/api/v1/documents/{str(doc_uuid)}/report?format=pdf")
    assert res_pdf.status_code == 200
    assert res_pdf.json()["status"] == "report_generated"
    assert "report_url" in res_pdf.json()
    assert "review_memorandum.pdf" in res_pdf.json()["report_url"] or "/download" in res_pdf.json()["report_url"]

    # Test run_report with format=docx
    res_docx = client.post(f"/api/v1/documents/{str(doc_uuid)}/report?format=docx")
    assert res_docx.status_code == 200
    assert res_docx.json()["status"] == "report_generated"
    assert "report_url" in res_docx.json()
    assert "review_memorandum.docx" in res_docx.json()["report_url"] or "/download" in res_docx.json()["report_url"]

    app.dependency_overrides.clear()

